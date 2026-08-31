# import pandas as pd
# import re,os
# from openpyxl import Workbook
# from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
# from openpyxl.utils import get_column_letter


# def missing_lines_coverage(OUTPUT_PATH,fs_df,missing_df,OUTPUT_FILE):
#     missing_content_coverage_report = os.path.join(OUTPUT_PATH,OUTPUT_FILE)
#     fs_df.columns      = fs_df.columns.str.strip()
#     missing_df.columns = missing_df.columns.str.strip()

#     for col in ['Class', 'Method', 'line_range', 'Spec_file']:
#         if col in fs_df.columns:
#             fs_df[col] = fs_df[col].astype(str).str.strip()

#     for col in ['Class', 'Method', 'Line_ID']:
#         if col in missing_df.columns:
#             missing_df[col] = missing_df[col].astype(str).str.strip()

#     # ── 2. Helpers ────────────────────────────────────────────────────────────────
#     def parse_line_numbers(range_str: str) -> set:
#         """f219_133-f219_135 → {133,134,135};  f219_139 → {139}"""
#         nums = re.findall(r'[Ff]\d+_(\d+)', str(range_str).strip())
#         if not nums:
#             return set()
#         ints = [int(n) for n in nums]
#         if len(ints) == 1:
#             return {ints[0]}
#         return set(range(min(ints), max(ints) + 1))

#     def parse_missing_line_id(line_id: str):
#         """F201_103 → ('f201_', 103)"""
#         m = re.match(r'[Ff](\d+)_(\d+)', str(line_id).strip())
#         if not m:
#             return ('', None)
#         return (f'f{m.group(1)}_', int(m.group(2)))

#     def get_file_prefix(range_str: str) -> str:
#         """f219_133-f219_135 → 'f219_'"""
#         m = re.match(r'([Ff]\d+_)', str(range_str).strip())
#         return m.group(1).lower() if m else ''

#     def min_max_range(range_strings: list) -> str:
#         """
#         Given multiple range/Line_ID strings for the same Class+Method,
#         collect ALL line numbers and return ONE range from global min to global max.

#         Examples
#         --------
#         ['f10_100-f10_150', 'f10_200-f10_250', 'f10_300'] → 'f10_100-f10_300'
#         ['f10_50']                                          → 'f10_50'
#         """
#         all_nums, prefix = [], ''
#         for r in range_strings:
#             nums = re.findall(r'[Ff]\d+_(\d+)', str(r).strip())
#             all_nums.extend(int(n) for n in nums)
#             if not prefix:
#                 prefix = get_file_prefix(r)
#         if not all_nums:
#             return ''
#         mn, mx = min(all_nums), max(all_nums)
#         return f'{prefix}{mn}' if mn == mx else f'{prefix}{mn}-{prefix}{mx}'

#     def norm(cls, mth):
#         return (str(cls).lower().strip(), str(mth).lower().strip())

#     # ── 3. Build lookup from Excel 2: (Class, Method) → {lines, grouped_range} ───
#     #      Multiple Line_IDs for same Class+Method are collapsed to min-max range.
#     uncovered_lookup = {}
#     for (cls, mth), grp in missing_df.groupby(['Class', 'Method']):
#         unc_lines, raw_ids = set(), []
#         for lid in grp['Line_ID'].tolist():
#             _, ln = parse_missing_line_id(lid)
#             if ln is not None:
#                 unc_lines.add(ln)
#             raw_ids.append(str(lid).strip())
#         uncovered_lookup[norm(cls, mth)] = {
#             'lines':         unc_lines,
#             # Collapse all Line_IDs to a single min–max range string
#             'grouped_range': min_max_range(raw_ids),
#         }

#     # ── 4. Build lookup from Excel 1: (Class, Method) → {lines, grouped_range, chunk} ──
#     #      Multiple line_range rows for same Class+Method are collapsed to min-max range.
#     covered_lookup = {}
#     for (cls, mth), grp in fs_df.groupby(['Class', 'Method']):
#         all_lines, raw_ranges = set(), []
#         for r in grp['line_range'].tolist():
#             all_lines.update(parse_line_numbers(r))
#             raw_ranges.append(str(r).strip())
#         covered_lookup[norm(cls, mth)] = {
#             'lines':         all_lines,
#             # Collapse all line_ranges to a single min–max range string
#             'grouped_range': min_max_range(raw_ranges),
#             'chunk_name':    ', '.join(sorted(grp['Spec_file'].unique().tolist())),
#             'orig_cls':      cls,
#             'orig_mth':      mth,
#         }

#     # ── 5. Classify — one output row per unique (Class, Method) ──────────────────
#     fully_covered_rows     = []
#     partially_covered_rows = []
#     completely_uncovered_rows = []

#     for nk, cov_info in covered_lookup.items():
#         cls   = cov_info['orig_cls']
#         mth   = cov_info['orig_mth']
#         chunk = cov_info['chunk_name']
#         cov_lines     = cov_info['lines']
#         cov_range_str = cov_info['grouped_range']   # min-max of Excel 1 ranges

#         if nk in uncovered_lookup:
#             unc_info  = uncovered_lookup[nk]
#             unc_lines = unc_info['lines']
#             unc_range_str = unc_info['grouped_range']   # min-max of Excel 2 Line_IDs

#             truly_covered = cov_lines - unc_lines

#             if not truly_covered:
#                 # All covered lines are also flagged missing → Completely Not Covered
#                 completely_uncovered_rows.append({
#                     'Class':      cls,
#                     'Method':     mth,
#                     'Line_Range': cov_range_str,
#                     'Chunk_Name': chunk,
#                 })
#             else:
#                 # Present in both sheets → Partially Covered
#                 # Covered_Line_Range   = grouped min-max of Excel 1 line_ranges
#                 # Uncovered_Line_Range = grouped min-max of Excel 2 Line_IDs
#                 partially_covered_rows.append({
#                     'Class':                cls,
#                     'Method':               mth,
#                     'Covered_Line_Range':   cov_range_str,
#                     'Uncovered_Line_Range': unc_range_str,
#                     'Chunk_Name':           chunk,
#                 })
#         else:
#             # Not in missing sheet at all → Fully Covered
#             fully_covered_rows.append({
#                 'Class':      cls,
#                 'Method':     mth,
#                 'Line_Range': cov_range_str,
#                 'Chunk_Name': chunk,
#             })

#     # ── 6. Completely Not Covered — entries ONLY in Excel 2 (not in Excel 1) ─────
#     fs_norm_keys = set(covered_lookup.keys())
#     for nk, info in uncovered_lookup.items():
#         if nk not in fs_norm_keys:
#             completely_uncovered_rows.append({
#                 'Class':      nk[0],
#                 'Method':     nk[1],
#                 'Line_Range': info['grouped_range'],
#                 'Chunk_Name': '',
#             })

#     # ── 7. Build DataFrames ───────────────────────────────────────────────────────
#     df_full    = pd.DataFrame(fully_covered_rows,
#                             columns=['Class', 'Method', 'Line_Range', 'Chunk_Name'])
#     df_partial = pd.DataFrame(partially_covered_rows,
#                             columns=['Class', 'Method', 'Covered_Line_Range',
#                                     'Uncovered_Line_Range', 'Chunk_Name'])
#     df_uncov   = pd.DataFrame(completely_uncovered_rows,
#                             columns=['Class', 'Method', 'Line_Range', 'Chunk_Name'])

#     print(f"Fully covered         : {len(df_full)} rows")
#     print(f"Partially covered     : {len(df_partial)} rows")
#     print(f"Completely not covered: {len(df_uncov)} rows")

#     # ── 8. Write Excel ────────────────────────────────────────────────────────────
    
#     wb = Workbook()

#     HEADER_FONT  = Font(name='Arial', bold=True, color='FFFFFF', size=11)
#     DATA_FONT    = Font(name='Arial', size=10)
#     ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
#     ALIGN_LEFT   = Alignment(horizontal='left',   vertical='center', wrap_text=True)
#     FILLS = {
#         'green':   PatternFill('solid', fgColor='1F7A4A'),
#         'orange':  PatternFill('solid', fgColor='C45911'),
#         'red':     PatternFill('solid', fgColor='A00000'),
#         'row_alt': PatternFill('solid', fgColor='F2F2F2'),
#     }
#     thin   = Side(style='thin', color='BBBBBB')
#     BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

#     def write_sheet(ws, headers, rows_df, fill_key):
#         ws.row_dimensions[1].height = 22
#         for ci, h in enumerate(headers, 1):
#             c = ws.cell(row=1, column=ci, value=h)
#             c.font = HEADER_FONT; c.fill = FILLS[fill_key]
#             c.alignment = ALIGN_CENTER; c.border = BORDER
#         for ri, row in enumerate(rows_df.itertuples(index=False), 2):
#             fill = FILLS['row_alt'] if ri % 2 == 0 else PatternFill()
#             for ci, val in enumerate(row, 1):
#                 c = ws.cell(row=ri, column=ci, value=val)
#                 c.font = DATA_FONT; c.fill = fill
#                 c.alignment = ALIGN_LEFT; c.border = BORDER
#             ws.row_dimensions[ri].height = 18
#         for ci, col_cells in enumerate(ws.columns, 1):
#             max_len = max((len(str(c.value or '')) for c in col_cells), default=10)
#             ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 4, 60)
#         ws.freeze_panes = 'A2'

#     ws1 = wb.active
#     ws1.title = 'Fully_Covered'
#     write_sheet(ws1, ['Class', 'Method', 'Line_Range', 'Chunk_Name'], df_full, 'green')

#     ws2 = wb.create_sheet('Partially_Covered')
#     write_sheet(ws2, ['Class', 'Method', 'Covered_Line_Range', 'Uncovered_Line_Range', 'Chunk_Name'],
#                 df_partial, 'orange')

#     ws3 = wb.create_sheet('Completely_Not_Covered')
#     write_sheet(ws3, ['Class', 'Method', 'Line_Range', 'Chunk_Name'], df_uncov, 'red')

#     wb.save(missing_content_coverage_report)
#     print(f"Saved → {missing_content_coverage_report}")
#     return missing_content_coverage_report

import pandas as pd
import re, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def missing_lines_coverage(OUTPUT_PATH,final_spec_parent_chunk_for_tech_spec,fs_df, missing_df, OUTPUT_FILE):
    """
    Parameters
    ----------
    fs_df, missing_df   : DataFrames (Excel 1 and Excel 2)
    OUTPUT_PATH         : folder to write the report into
    OUTPUT_FILE         : filename for the report
    final_spec_parent_chunk_for_tech_spec : path to the folder that contains
                          chunk files (e.g. chunk_001.txt, chunk_42.md …).
                          Each file name or its content is expected to contain
                          line-ID tokens like f219_133 so we can tell which
                          line numbers that chunk covers.
                          Pass '' or None to skip this lookup.
    """
    missing_content_coverage_report = os.path.join(OUTPUT_PATH, OUTPUT_FILE)
    # Accept either a file path (str) or an already-loaded DataFrame
    if isinstance(fs_df, str):
        fs_df = pd.read_excel(fs_df, sheet_name='FS Methods')
    if isinstance(missing_df, str):
        missing_df = pd.read_excel(missing_df, sheet_name='Missing_Lines')
    fs_df.columns      = fs_df.columns.str.strip()
    missing_df.columns = missing_df.columns.str.strip()

    for col in ['Class', 'Method', 'line_range', 'Spec_file']:
        if col in fs_df.columns:
            fs_df[col] = fs_df[col].astype(str).str.strip()

    for col in ['Class', 'Method', 'Line_ID']:
        if col in missing_df.columns:
            missing_df[col] = missing_df[col].astype(str).str.strip()

    # ── 2. Helpers ────────────────────────────────────────────────────────────

    def parse_line_numbers(range_str: str) -> set:
        """f219_133-f219_135 → {133,134,135};  f219_139 → {139}"""
        nums = re.findall(r'[Ff]\d+_(\d+)', str(range_str).strip())
        if not nums:
            return set()
        ints = [int(n) for n in nums]
        if len(ints) == 1:
            return {ints[0]}
        return set(range(min(ints), max(ints) + 1))

    def parse_missing_line_id(line_id: str):
        """F201_103 → ('f201_', 103)"""
        m = re.match(r'[Ff](\d+)_(\d+)', str(line_id).strip())
        if not m:
            return ('', None)
        return (f'f{m.group(1)}_', int(m.group(2)))

    def get_file_prefix(range_str: str) -> str:
        """f219_133-f219_135 → 'f219_'"""
        m = re.match(r'([Ff]\d+_)', str(range_str).strip())
        return m.group(1).lower() if m else ''

    def min_max_range(range_strings: list) -> str:
        """
        Given multiple range/Line_ID strings for the same Class+Method,
        collect ALL line numbers and return ONE range from global min to global max.

        Examples
        --------
        ['f10_100-f10_150', 'f10_200-f10_250', 'f10_300'] → 'f10_100-f10_300'
        ['f10_50']                                          → 'f10_50'
        """
        all_nums, prefix = [], ''
        for r in range_strings:
            nums = re.findall(r'[Ff]\d+_(\d+)', str(r).strip())
            all_nums.extend(int(n) for n in nums)
            if not prefix:
                prefix = get_file_prefix(r)
        if not all_nums:
            return ''
        mn, mx = min(all_nums), max(all_nums)
        return f'{prefix}{mn}' if mn == mx else f'{prefix}{mn}-{prefix}{mx}'

    def norm(cls, mth):
        return (str(cls).lower().strip(), str(mth).lower().strip())

    # ── 2b. Build chunk-folder index ─────────────────────────────────────────
    # chunk_line_index  : { line_number(int) → [chunk_file_name, …] }
    # Strategy: scan every file in the folder; collect all fXXX_NNN tokens
    # found in the *file name* first, then fall back to scanning file *content*.
    # Both approaches are tried so it works regardless of naming convention.

    chunk_line_index = {}   # line_num(int) → [chunk file names]

    def _index_chunk_folder(folder):
        """Populate chunk_line_index from files inside `folder`."""
        if not folder or not os.path.isdir(folder):
            return

        LINE_ID_RE = re.compile(r'[Ff]\d+_(\d+)')

        for fname in os.listdir(folder):
            fpath = os.path.join(folder, fname)
            if not os.path.isfile(fpath):
                continue

            line_nums = set()

            # ① Try to extract line numbers from the file NAME itself
            nums_in_name = LINE_ID_RE.findall(fname)
            if nums_in_name:
                ints = [int(n) for n in nums_in_name]
                # If the name encodes a range (two numbers) expand it
                if len(ints) >= 2:
                    line_nums.update(range(min(ints), max(ints) + 1))
                else:
                    line_nums.update(ints)

            # ② Fall back: scan file content for fXXX_NNN tokens
            if not line_nums:
                try:
                    with open(fpath, 'r', encoding='utf-8', errors='ignore') as fh:
                        content = fh.read()
                    for n in LINE_ID_RE.findall(content):
                        line_nums.add(int(n))
                except Exception:
                    pass  # binary or unreadable — skip

            # Register every discovered line number → this chunk file
            for ln in line_nums:
                chunk_line_index.setdefault(ln, []).append(fname)

    _index_chunk_folder(final_spec_parent_chunk_for_tech_spec)

    def find_chunk_for_range(line_range_str):
        """
        Given a grouped range string like 'f10_100-f10_150', expand it to
        individual line numbers, look each up in chunk_line_index, and return
        a deduplicated, sorted, comma-joined string of matching chunk names.
        Returns '' if the index is empty or nothing matched.
        """
        if not chunk_line_index:
            return ''

        target_lines = parse_line_numbers(line_range_str)
        if not target_lines:
            return ''

        matched_chunks = set()
        for ln in target_lines:
            for cname in chunk_line_index.get(ln, []):
                matched_chunks.add(cname)

        return ', '.join(sorted(matched_chunks))

    # ── 3. Build lookup from Excel 2: (Class, Method) → {lines, grouped_range} ──
    uncovered_lookup = {}
    for (cls, mth), grp in missing_df.groupby(['Class', 'Method']):
        unc_lines, raw_ids = set(), []
        for lid in grp['Line_ID'].tolist():
            _, ln = parse_missing_line_id(lid)
            if ln is not None:
                unc_lines.add(ln)
            raw_ids.append(str(lid).strip())
        uncovered_lookup[norm(cls, mth)] = {
            'lines':         unc_lines,
            'grouped_range': min_max_range(raw_ids),
        }

    # ── 4. Build lookup from Excel 1: (Class, Method) → {lines, grouped_range, chunk} ──
    covered_lookup = {}
    for (cls, mth), grp in fs_df.groupby(['Class', 'Method']):
        all_lines, raw_ranges = set(), []
        for r in grp['line_range'].tolist():
            all_lines.update(parse_line_numbers(r))
            raw_ranges.append(str(r).strip())
        covered_lookup[norm(cls, mth)] = {
            'lines':         all_lines,
            'grouped_range': min_max_range(raw_ranges),
            'chunk_name':    ', '.join(sorted(grp['Spec_file'].unique().tolist())),
            'orig_cls':      cls,
            'orig_mth':      mth,
        }

    # ── 5. Classify — one output row per unique (Class, Method) ──────────────
    fully_covered_rows        = []
    partially_covered_rows    = []
    completely_uncovered_rows = []

    for nk, cov_info in covered_lookup.items():
        cls   = cov_info['orig_cls']
        mth   = cov_info['orig_mth']
        chunk = cov_info['chunk_name']
        cov_lines     = cov_info['lines']
        cov_range_str = cov_info['grouped_range']

        if nk in uncovered_lookup:
            unc_info      = uncovered_lookup[nk]
            unc_lines     = unc_info['lines']
            unc_range_str = unc_info['grouped_range']
            truly_covered = cov_lines - unc_lines

            if not truly_covered:
                # All covered lines are also flagged missing → Completely Not Covered
                # chunk is already known from Excel 1
                completely_uncovered_rows.append({
                    'Class':      cls,
                    'Method':     mth,
                    'Line_Range': cov_range_str,
                    'Chunk_Name': chunk,
                })
            else:
                partially_covered_rows.append({
                    'Class':                cls,
                    'Method':               mth,
                    'Covered_Line_Range':   cov_range_str,
                    'Uncovered_Line_Range': unc_range_str,
                    'Chunk_Name':           chunk,
                })
        else:
            fully_covered_rows.append({
                'Class':      cls,
                'Method':     mth,
                'Line_Range': cov_range_str,
                'Chunk_Name': chunk,
            })

    # ── 6. Completely Not Covered — entries ONLY in Excel 2 (not in Excel 1) ─
    #      These had no Chunk_Name before. Now we look it up from the chunk folder.
    fs_norm_keys = set(covered_lookup.keys())
    for nk, info in uncovered_lookup.items():
        if nk not in fs_norm_keys:
            line_range_str = info['grouped_range']

            # ▶ NEW: resolve Chunk_Name by scanning the chunk folder
            chunk_name = find_chunk_for_range(line_range_str)

            completely_uncovered_rows.append({
                'Class':      nk[0],
                'Method':     nk[1],
                'Line_Range': line_range_str,
                'Chunk_Name': chunk_name,   # populated from folder scan
            })

    # ── 7. Build DataFrames ───────────────────────────────────────────────────
    df_full    = pd.DataFrame(fully_covered_rows,
                              columns=['Class', 'Method', 'Line_Range', 'Chunk_Name'])
    df_partial = pd.DataFrame(partially_covered_rows,
                              columns=['Class', 'Method', 'Covered_Line_Range',
                                       'Uncovered_Line_Range', 'Chunk_Name'])
    df_uncov   = pd.DataFrame(completely_uncovered_rows,
                              columns=['Class', 'Method', 'Line_Range', 'Chunk_Name'])

    print(f"Fully covered         : {len(df_full)} rows")
    print(f"Partially covered     : {len(df_partial)} rows")
    print(f"Completely not covered: {len(df_uncov)} rows")

    # ── 8. Write Excel ────────────────────────────────────────────────────────
    wb = Workbook()

    HEADER_FONT  = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    DATA_FONT    = Font(name='Arial', size=10)
    ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ALIGN_LEFT   = Alignment(horizontal='left',   vertical='center', wrap_text=True)
    FILLS = {
        'green':   PatternFill('solid', fgColor='1F7A4A'),
        'orange':  PatternFill('solid', fgColor='C45911'),
        'red':     PatternFill('solid', fgColor='A00000'),
        'row_alt': PatternFill('solid', fgColor='F2F2F2'),
    }
    thin   = Side(style='thin', color='BBBBBB')
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    def write_sheet(ws, headers, rows_df, fill_key):
        ws.row_dimensions[1].height = 22
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = HEADER_FONT; c.fill = FILLS[fill_key]
            c.alignment = ALIGN_CENTER; c.border = BORDER
        for ri, row in enumerate(rows_df.itertuples(index=False), 2):
            fill = FILLS['row_alt'] if ri % 2 == 0 else PatternFill()
            for ci, val in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.font = DATA_FONT; c.fill = fill
                c.alignment = ALIGN_LEFT; c.border = BORDER
            ws.row_dimensions[ri].height = 18
        for ci, col_cells in enumerate(ws.columns, 1):
            max_len = max((len(str(c.value or '')) for c in col_cells), default=10)
            ws.column_dimensions[get_column_letter(ci)].width = min(max_len + 4, 60)
        ws.freeze_panes = 'A2'

    ws1 = wb.active
    ws1.title = 'Fully_Covered'
    write_sheet(ws1, ['Class', 'Method', 'Line_Range', 'Chunk_Name'], df_full, 'green')

    ws2 = wb.create_sheet('Partially_Covered')
    write_sheet(ws2, ['Class', 'Method', 'Covered_Line_Range', 'Uncovered_Line_Range', 'Chunk_Name'],
                df_partial, 'orange')

    ws3 = wb.create_sheet('Completely_Not_Covered')
    write_sheet(ws3, ['Class', 'Method', 'Line_Range', 'Chunk_Name'], df_uncov, 'red')

    wb.save(missing_content_coverage_report)
    print(f"Saved → {missing_content_coverage_report}")
    return missing_content_coverage_report

# OUTPUT_PATH = r"C:\Users\Barath\Desktop\Reverse_Engineering_rabobank\output"
# final_spec_parent_chunk_for_tech_spec = r"C:\Project\Reverse_Engineering_framework\end_to_end_framework_result_all_application\day_trade\day_trade\testing_28_07_copilot\chunk_for_3_logics"
# fs_df = r"C:\Users\Barath\Desktop\Reverse_Engineering_rabobank\output\fs_methods.xlsx"
# missing_df = r"C:\Project\Reverse_Engineering_framework\end_to_end_framework_result_all_application\day_trade\day_trade\test_29_07\Validation_Reports\recursive_chunk_vs_spec_line_number_validation.xlsx"
# OUTPUT_FILE = "missing_coverage_report.xlsx"
# missing_lines_coverage(OUTPUT_PATH,final_spec_parent_chunk_for_tech_spec,fs_df, missing_df, OUTPUT_FILE)